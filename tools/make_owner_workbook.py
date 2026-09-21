# -*- coding: utf-8 -*-
"""Build the owner/deadline workbook that Adriaan and Jan edit in Excel Online.

The page is static and rebuilt every Monday, so an edit made in the browser
would live in one person's browser and be gone by Tuesday. The two fields that
are genuinely a person's call - WHO owns an action and BY WHEN - therefore live
outside the page, in one workbook on SharePoint, one row per action.

Nothing else belongs in it. Status, counts and closure are computed by
tools/eval_conditions.py from the stored conditions and are never read from
the workbook: if the workbook could set status, a person could mark an item
done that the data says is not.

    python3 tools/make_owner_workbook.py         # write build/action_owners.xlsx
"""
import datetime, os, re, sys, zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import render_actions as RA                                    # noqa: E402
from openpyxl import Workbook                                  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation    # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill        # noqa: E402
from openpyxl.utils import get_column_letter                    # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "build", "action_owners.xlsx")
OWNERS = ["Adriaan Mol", "James Walker", "Jan de Graaf",
          "Angelo Nahavitatsara", "Lanja Randriamanantena", "MadAvance",
          "Angelo Nahavitatsara / MadAvance",
          "Lanja Randriamanantena / MadAvance",
          "MadAvance field teams", "Jan / Endur'O team", "Coddy", "Cathy",
          "Earthood", "Unassigned"]


from xlsx_invariants import is_opc_part, verify      # noqa: E402,F401


def main():
    idx = open(os.path.join(REPO, "index.html"), encoding="utf8").read()
    acts = sorted(RA.actions(idx), key=lambda a: a["id"])
    wb = Workbook()
    ws = wb.active
    ws.title = "actions"
    head = ["id", "owner", "deadline", "what it is (read-only)"]
    ws.append(head)
    for i, h in enumerate(head, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EDF1")
    # Cells are written explicitly, not by append(), and a blank deadline is
    # left genuinely EMPTY. append() with None produced <c r="C106" s="3"
    # t="n"></c> - an explicit numeric type with no value - on every row with
    # no date, which is malformed enough for Excel Online to "repair" the
    # workbook, and a repaired workbook always opens read-only.
    #
    # Real deadlines are written as datetime.date with a yyyy-mm-dd number
    # format, so they are dates rather than text and the date validation on
    # column C actually guards something.
    for i, a in enumerate(acts, start=2):
        ws.cell(row=i, column=1, value=a["id"])
        ws.cell(row=i, column=2, value=a["owner"])
        c = ws.cell(row=i, column=3)
        if a["due"]:
            c.value = a["due"]
            c.number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=4, value=a["title"][:120])
    for col, w in zip("ABCD", (44, 34, 14, 90)):
        ws.column_dimensions[col].width = w
    last = ws.max_row
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(OWNERS),
                        allow_blank=True, showDropDown=False)
    dv.error = "Pick an owner from the list, or add them to tools/make_owner_workbook.py"
    ws.add_data_validation(dv)
    dv.add(f"B2:B{last}")
    dd = DataValidation(type="date", operator="greaterThan",
                        formula1="DATE(2024,1,1)", allow_blank=True)
    dd.error = "Deadlines are dates. Leave blank if no date is set yet."
    ws.add_data_validation(dd)
    dd.add(f"C2:C{last}")
    for r in range(2, last + 1):
        ws.cell(row=r, column=1).font = Font(color="7B858D")
        ws.cell(row=r, column=4).font = Font(color="7B858D")
        ws.cell(row=r, column=4).alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "A2"

    note = wb.create_sheet("read me")
    for line in [
        "SaniTap weekly water report - owner and deadline",
        "",
        "One row per action in the report. Edit column B (owner) and column C",
        "(deadline) only. The build reads this file each Monday and renders",
        "those two fields from it.",
        "",
        "Column A is the action id and must not be changed: it is how a row is",
        "matched to the report. Column D is a copy of the title, for reading",
        "only - editing it changes nothing.",
        "",
        "Status, the counts, and whether an item is open or closed are NOT in",
        "this workbook and cannot be set here. They are computed from stored",
        "closing conditions every build, so an item closes when the data says",
        "it is done and reopens if that stops being true.",
        "",
        "If the build cannot read this file, the report falls back to the",
        "owners and deadlines written in the page and says so on the page.",
        "",
        f"Seeded {datetime.date.today().isoformat()} from "
        f"{len(acts)} actions by tools/make_owner_workbook.py",
    ]:
        # a blank line must be a blank ROW, not a cell holding "" - an empty
        # inlineStr body is the same malformation as the typed-empty deadline
        if line:
            note.append([line])
        else:
            note.append([])
    note.column_dimensions["A"].width = 78

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    verify(OUT, expect_ids=[a["id"] for a in acts], expect_owners=set(OWNERS))
    print(f"{OUT}: {len(acts)} rows, {len(OWNERS)} owners in the dropdown")
    return 0


if __name__ == "__main__":
    sys.exit(main())
